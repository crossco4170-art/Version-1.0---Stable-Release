from pathlib import Path

from docx import Document


def create_sample_docx(path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Jane Doe").bold = True
    paragraph.add_run("\n(555) 123-4567")
    paragraph.add_run("\njane@example.com")
    paragraph.add_run("\nExperience: 5 years in software development")
    paragraph.add_run("\nSkills: Python, SQL, FastAPI")
    paragraph.add_run("\nEducation: BS Computer Science")
    paragraph.add_run("\nEmployment History: Senior Engineer at Acme Corp")
    document.save(path)


def create_sample_pdf(path: Path) -> None:
    stream_text = b"BT\n/F1 24 Tf\n72 720 Td\n(Jane Doe) Tj\nET\n"
    objects = [
        (1, b"<< /Type /Catalog /Pages 2 0 R >>"),
        (2, b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>"),
        (
            3,
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        ),
        (4, b"<< /Length " + str(len(stream_text)).encode() + b" >>\nstream\n" + stream_text + b"\nendstream"),
        (5, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"),
    ]

    pdf_bytes = b"%PDF-1.4\n"
    offsets = []
    for object_number, body in objects:
        offsets.append(len(pdf_bytes))
        pdf_bytes += f"{object_number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"

    xref_offset = len(pdf_bytes)
    pdf_bytes += b"xref\n0 " + str(len(objects) + 1).encode("ascii") + b"\n"
    pdf_bytes += b"0000000000 65535 f \n"
    for offset in offsets:
        pdf_bytes += f"{offset:010d} 00000 n \n".encode("ascii")

    pdf_bytes += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii")
    path.write_bytes(pdf_bytes)


if __name__ == "__main__":
    base_dir = Path(__file__).resolve().parent
    create_sample_docx(base_dir / "sample_resume.docx")
    create_sample_pdf(base_dir / "sample_resume.pdf")
