from __future__ import annotations

import csv
from calendar import monthrange
from datetime import date, datetime, time, timedelta
from pathlib import Path
from statistics import mean
from typing import Any

from sqlalchemy.orm import Session

from config.settings import get_settings
from database.connection import get_session
from models.applicant import Applicant
from models.interview import Interview
from utils.datetime_policy import utc_now_naive
from utils.exceptions import BlackcrestDependencyError, BlackcrestInputError


class ReportingService:
    """Generate recruiting reports and export them to common business formats."""

    def __init__(self, session: Session | None = None, database_url: str | None = None) -> None:
        settings = get_settings(validate_required=False)
        resolved_url = database_url or settings.database_url
        self.session = session or get_session(database_url=resolved_url, ensure_schema=True)

        self.score_threshold = settings.score_notification_threshold

    def generate_daily_report(self, target_date: date | None = None) -> dict[str, Any]:
        report_date = target_date or date.today()
        return self._build_period_report(
            report_type="daily",
            period_start=report_date,
            period_end=report_date,
        )

    def generate_weekly_report(self, target_date: date | None = None) -> dict[str, Any]:
        anchor = target_date or date.today()
        period_start = anchor - timedelta(days=anchor.weekday())
        period_end = period_start + timedelta(days=6)
        return self._build_period_report(
            report_type="weekly",
            period_start=period_start,
            period_end=period_end,
        )

    def generate_monthly_report(self, year: int | None = None, month: int | None = None) -> dict[str, Any]:
        today = date.today()
        resolved_year = year or today.year
        resolved_month = month or today.month
        if resolved_month < 1 or resolved_month > 12:
            raise BlackcrestInputError("month must be between 1 and 12")

        first_day = date(resolved_year, resolved_month, 1)
        last_day = date(resolved_year, resolved_month, monthrange(resolved_year, resolved_month)[1])
        return self._build_period_report(
            report_type="monthly",
            period_start=first_day,
            period_end=last_day,
        )

    def generate_applicant_pipeline_report(self) -> dict[str, Any]:
        applicants = list(self.session.query(Applicant).all())
        pipeline = self._build_pipeline_rows(applicants)
        scores = [item.score for item in applicants if item.score is not None]

        return {
            "report_type": "applicant_pipeline",
            "generated_at": utc_now_naive().isoformat(timespec="seconds"),
            "summary": {
                "total_applicants": len(applicants),
                "average_score": round(mean(scores), 2) if scores else 0.0,
                "above_threshold": len([score for score in scores if score >= self.score_threshold]),
            },
            "pipeline": pipeline,
            "applicants": self._build_applicant_rows(applicants),
        }

    def export_report_to_csv(self, report: dict[str, Any], output_path: str | Path) -> Path:
        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        summary = report.get("summary", {})
        pipeline = report.get("pipeline", [])
        applicants = report.get("applicants", [])

        with destination.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["Report Type", report.get("report_type", "unknown")])
            writer.writerow(["Generated At", report.get("generated_at", "")])
            writer.writerow([])

            writer.writerow(["Summary"])
            writer.writerow(["Metric", "Value"])
            for key, value in summary.items():
                writer.writerow([key, value])
            writer.writerow([])

            writer.writerow(["Applicant Pipeline"])
            writer.writerow(["Status", "Count"])
            for row in pipeline:
                writer.writerow([row.get("status", ""), row.get("count", 0)])
            writer.writerow([])

            writer.writerow(["Applicants"])
            writer.writerow(["ID", "Name", "Status", "Score", "Email", "Date Added"])
            for row in applicants:
                writer.writerow(
                    [
                        row.get("id", ""),
                        row.get("name", ""),
                        row.get("current_status", ""),
                        row.get("score", ""),
                        row.get("email", ""),
                        row.get("date_added", ""),
                    ]
                )
        return destination

    def export_report_to_excel(self, report: dict[str, Any], output_path: str | Path) -> Path:
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise BlackcrestDependencyError(
                "openpyxl is required for Excel export. Install it from requirements.txt."
            ) from exc

        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()

        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        summary_sheet.append(["Report Type", report.get("report_type", "unknown")])
        summary_sheet.append(["Generated At", report.get("generated_at", "")])
        summary_sheet.append([])
        summary_sheet.append(["Metric", "Value"])
        for key, value in report.get("summary", {}).items():
            summary_sheet.append([key, value])

        pipeline_sheet = workbook.create_sheet(title="Pipeline")
        pipeline_sheet.append(["Status", "Count"])
        for row in report.get("pipeline", []):
            pipeline_sheet.append([row.get("status", ""), row.get("count", 0)])

        applicants_sheet = workbook.create_sheet(title="Applicants")
        applicants_sheet.append(["ID", "Name", "Status", "Score", "Email", "Date Added"])
        for row in report.get("applicants", []):
            applicants_sheet.append(
                [
                    row.get("id", ""),
                    row.get("name", ""),
                    row.get("current_status", ""),
                    row.get("score", ""),
                    row.get("email", ""),
                    row.get("date_added", ""),
                ]
            )

        workbook.save(destination)
        return destination

    def export_report_to_pdf(self, report: dict[str, Any], output_path: str | Path) -> Path:
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise BlackcrestDependencyError(
                "reportlab is required for PDF export. Install it from requirements.txt."
            ) from exc

        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        pdf = canvas.Canvas(str(destination), pagesize=letter)
        width, height = letter
        cursor_y = height - 50

        def write_line(text: str, gap: int = 16) -> None:
            nonlocal cursor_y
            if cursor_y < 60:
                pdf.showPage()
                cursor_y = height - 50
            pdf.drawString(40, cursor_y, text)
            cursor_y -= gap

        write_line(f"Report Type: {report.get('report_type', 'unknown')}")
        write_line(f"Generated At: {report.get('generated_at', '')}")
        write_line("")

        write_line("Summary")
        for key, value in report.get("summary", {}).items():
            write_line(f"- {key}: {value}")

        write_line("")
        write_line("Applicant Pipeline")
        for row in report.get("pipeline", []):
            write_line(f"- {row.get('status', '')}: {row.get('count', 0)}")

        write_line("")
        write_line("Applicants")
        for row in report.get("applicants", []):
            write_line(
                f"- #{row.get('id', '')} {row.get('name', '')} | "
                f"{row.get('current_status', '')} | "
                f"Score: {row.get('score', '')}"
            )

        pdf.save()
        return destination

    def _build_period_report(
        self,
        *,
        report_type: str,
        period_start: date,
        period_end: date,
    ) -> dict[str, Any]:
        start_dt = datetime.combine(period_start, time.min)
        end_dt = datetime.combine(period_end, time.max)

        applicants = list(
            self.session.query(Applicant)
            .filter(Applicant.date_added >= start_dt, Applicant.date_added <= end_dt)
            .all()
        )

        interviews_scheduled = (
            self.session.query(Interview)
            .filter(Interview.interview_date.is_not(None))
            .filter(Interview.interview_date >= start_dt, Interview.interview_date <= end_dt)
            .count()
        )

        scores = [item.score for item in applicants if item.score is not None]
        return {
            "report_type": report_type,
            "generated_at": utc_now_naive().isoformat(timespec="seconds"),
            "period_start": period_start.isoformat(),
            "period_end": period_end.isoformat(),
            "summary": {
                "total_new_applicants": len(applicants),
                "interviews_scheduled": interviews_scheduled,
                "average_score": round(mean(scores), 2) if scores else 0.0,
                "above_threshold": len([score for score in scores if score >= self.score_threshold]),
            },
            "pipeline": self._build_pipeline_rows(applicants),
            "applicants": self._build_applicant_rows(applicants),
        }

    def _build_pipeline_rows(self, applicants: list[Applicant]) -> list[dict[str, Any]]:
        counts: dict[str, int] = {}
        for applicant in applicants:
            status = (applicant.current_status or "unassigned").strip().lower()
            counts[status] = counts.get(status, 0) + 1

        rows = [{"status": status.title(), "count": count} for status, count in counts.items()]
        return sorted(rows, key=lambda item: item["status"])

    def _build_applicant_rows(self, applicants: list[Applicant]) -> list[dict[str, Any]]:
        sorted_applicants = sorted(applicants, key=lambda item: item.date_added or datetime.min)
        rows: list[dict[str, Any]] = []
        for applicant in sorted_applicants:
            rows.append(
                {
                    "id": applicant.id,
                    "name": applicant.name,
                    "current_status": (applicant.current_status or "unassigned").title(),
                    "score": applicant.score if applicant.score is not None else "",
                    "email": applicant.email or "",
                    "date_added": applicant.date_added.strftime("%Y-%m-%d") if applicant.date_added else "",
                }
            )
        return rows
